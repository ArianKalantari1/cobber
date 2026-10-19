#!/usr/bin/env python3
"""Import scraped city bar menus (cleaned .xlsx) into a frontier corpus.

Reads the cocktail sheet of each cleaned city dataset (Sydney, Paris, London,
...) and writes one row per drink to ``data/raw/city_bar_cocktails.json``:
city, venue, drink name, base spirit(s) and the listed ingredients as a list,
plus the descriptive columns when present. Prices are dropped: they date fast
and don't feed Cobber.

These are working bars' signatures, so ``normalize.py`` treats them as a
*frontier* corpus (validated-novel evidence with venue attribution), never as
canon. Re-importing a dataset replaces its earlier rows; rows from other
datasets already in the JSON are kept, so new cities can be added one file at
a time without the original spreadsheets.

The city comes from the sheet's ``City`` column, else ``--city``, else the
file-name prefix when it is a known city. Anything else is refused rather than
guessed. Menus in another language (Milan) can be put into English with
``--translate it``, using the hand-edited ``data/menu_glossary.json``; the
original wording is kept in ``ingredients_original``. Needs ``openpyxl``
(build-time only).

    python3 scripts/import_city_menus.py Sydney_Cocktail_Dataset_clean.xlsx ...
    python3 scripts/import_city_menus.py --translate it Milan_Cocktail_Intelligence.xlsx
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "data" / "raw" / "city_bar_cocktails.json"

# File-name prefixes accepted as a city without --city ("New_York_..." reads
# as "newyork"). Extend as datasets arrive; an unknown prefix (e.g.
# "Martini_...") must be named explicitly.
KNOWN_CITIES = {
    "sydney": "Sydney", "melbourne": "Melbourne", "paris": "Paris",
    "london": "London", "newyork": "New York", "milan": "Milan",
    "tokyo": "Tokyo", "singapore": "Singapore", "hongkong": "Hong Kong",
}

# Output field -> accepted column headers (first match wins). Two dataset
# layouts so far: the original scans ("All Listed Ingredients") and the
# New York / Milan scans ("Ingredients", "Ingredients Complete").
COLUMNS = {
    "city": ("City",),
    "venue": ("Venue", "Bar"),
    "venue_type": ("Venue Type",),
    "name": ("Cocktail",),
    "section": ("Menu / Section", "Menu Section"),
    "base": ("Base Spirit(s)",),
    "ingredients": ("All Listed Ingredients", "Ingredients"),
    "ingredients_complete": ("Ingredients Complete",),
    "menu_role": ("Popularity / Menu Role", "Menu Role"),
    "style": ("Style / Genre",),
    "technique": ("Technique / Service", "Technique"),
    "culinary_tags": ("Culinary Tags",),
    "source_url": ("Source URL", "Primary Menu URL"),
    "accessed": ("Access Date", "Source Date"),
}
REQUIRED = ("venue", "name", "ingredients")
GLOSSARY_PATH = ROOT / "data" / "menu_glossary.json"


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if hasattr(value, "date") and callable(value.date):
        return value.date().isoformat()
    text = str(value).strip()
    return "" if text.lower() in ("nan", "none", "n/a", "na", "-", "—") else text


def _split_ingredients(text: str) -> list[str]:
    separator = ";" if ";" in text else ","
    return [part.strip() for part in text.split(separator) if part.strip()]


def _city_from_filename(path: Path) -> str | None:
    stem = re.sub(r"[^a-z]", "", path.stem.lower())
    for key, city in KNOWN_CITIES.items():
        if stem.startswith(key):
            return city
    return None


Glossary = tuple[re.Pattern, dict[str, str]]


def load_glossary(language: str) -> Glossary:
    """One whole-word, case-insensitive pattern for a language's phrases.

    Longest phrases come first and every match is replaced exactly once, so a
    protected name ("punt e mes") is never re-translated word by word.
    """
    glossary = json.loads(GLOSSARY_PATH.read_text(encoding="utf-8")).get(language)
    if not glossary:
        raise SystemExit(f"No '{language}' glossary in {GLOSSARY_PATH.name}.")
    phrases = sorted(glossary, key=len, reverse=True)
    pattern = re.compile(r"(?<![\w'])(" + "|".join(map(re.escape, phrases)) + r")(?!\w)",
                         re.IGNORECASE)
    return pattern, {phrase.lower(): english for phrase, english in glossary.items()}


def _translate(item: str, glossary: Glossary) -> str:
    pattern, mapping = glossary
    item = pattern.sub(lambda match: mapping[match.group(0).lower()], item)
    return re.sub(r"\s{2,}", " ", item).strip()


def _find_sheet(workbook):
    for sheet in workbook.worksheets:
        header = [_cell_text(cell) for cell in next(sheet.iter_rows(max_row=1, values_only=True), ())]
        if "Cocktail" in header and any(name in header for name in COLUMNS["ingredients"]):
            return sheet, header
    return None, []


def read_dataset(path: Path, city_override: str | None,
                 rules: Glossary | None = None) -> list[dict]:
    """Rows from one cleaned workbook, in output shape."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet, header = _find_sheet(workbook)
    if sheet is None:
        raise SystemExit(f"{path.name}: no sheet with a 'Cocktail' and an ingredients column.")
    index = {}
    for field, names in COLUMNS.items():
        for name in names:
            if name in header:
                index[field] = header.index(name)
                break
    missing = [field for field in REQUIRED if field not in index]
    if missing:
        raise SystemExit(f"{path.name}: missing column(s) for {', '.join(missing)}.")
    fallback_city = city_override or _city_from_filename(path)
    if "city" not in index and not fallback_city:
        raise SystemExit(f"{path.name}: no City column and no known city in the file name; pass --city.")

    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        record = {field: _cell_text(values[col]) if col < len(values) else ""
                  for field, col in index.items()}
        if not record.get("name") or not record.get("ingredients"):
            continue
        record["city"] = record.get("city") or fallback_city
        record["ingredients"] = _split_ingredients(record["ingredients"])
        if rules:
            translated = [_translate(item, rules) for item in record["ingredients"]]
            if translated != record["ingredients"]:
                record["ingredients_original"] = record["ingredients"]
                record["ingredients"] = translated
        record["dataset"] = path.stem
        rows.append({key: value for key, value in record.items() if value})
    workbook.close()
    return rows


def _row_key(row: dict) -> str:
    return re.sub(r"[^a-z0-9]", "", f"{row.get('city')}|{row.get('venue')}|{row['name']}".lower())


def main() -> None:
    parser = argparse.ArgumentParser(description="Import cleaned city bar menu datasets (.xlsx).")
    parser.add_argument("files", nargs="+", type=Path, help="Cleaned dataset workbooks.")
    parser.add_argument("--city", help="City for every row of a dataset that has no City column.")
    parser.add_argument("--translate", metavar="LANG",
                        help="Translate ingredients to English with data/menu_glossary.json "
                             "(e.g. 'it' for Milan); originals are kept in ingredients_original.")
    args = parser.parse_args()

    rules = load_glossary(args.translate) if args.translate else None
    imported = {path.stem: read_dataset(path, args.city, rules) for path in args.files}

    existing = []
    if OUT_PATH.exists():
        existing = [row for row in json.loads(OUT_PATH.read_text(encoding="utf-8"))
                    if row.get("dataset") not in imported]

    merged, seen, duplicates = [], set(), 0
    for row in [*existing, *(row for rows in imported.values() for row in rows)]:
        key = _row_key(row)
        if key in seen:
            duplicates += 1
            continue
        seen.add(key)
        merged.append(row)
    merged.sort(key=lambda row: (row.get("city", ""), row.get("venue", ""), row["name"]))

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with OUT_PATH.open("w", encoding="utf-8") as handle:
        json.dump(merged, handle, indent=2, ensure_ascii=False, sort_keys=True)
        handle.write("\n")

    for stem, rows in imported.items():
        print(f"{stem}: {len(rows)} drinks")
    cities = Counter(row.get("city", "?") for row in merged)
    print(f"Wrote {len(merged)} drinks to {OUT_PATH} ({dict(sorted(cities.items()))}); "
          f"skipped {duplicates} duplicate(s).")


if __name__ == "__main__":
    main()
